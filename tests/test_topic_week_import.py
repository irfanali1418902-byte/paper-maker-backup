"""topic_week_import_service — Excel se hafta-war plan. R7 Marhala 2."""

import io

import openpyxl
import pytest

from app.repositories import syllabus_repository, topic_week_plan_repository
from app.services import topic_week_import_service


def _topic(topic_id: str, title: str, unit_no: int = 1,
           subject: str = "Mathematics", grade: str = "Pre Year 2") -> None:
    syllabus_repository.insert(
        topic_id=topic_id,
        subject=subject,
        grade=grade,
        unit_no=unit_no,
        unit_title="Mathematics",
        page_range="",
        subtopic_title=title,
        activity_type="Introduction",
        page_no=None,
        learning_outcome="",
    )


def _sheet(rows, header=None) -> bytes:
    """rows: [[topic_id, title, unit_no, week_no], ...] — template ka shape."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header or topic_week_import_service.TEMPLATE_HEADER)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _weeks(subject="Mathematics", grade="Pre Year 2") -> dict:
    return {
        r["syllabus_topic_id"]: r["week_no"]
        for r in topic_week_plan_repository.list_resolved(subject, grade)
    }


# ---- template ----


def test_template_lists_every_topic_including_unplanned(test_db):
    _topic("t1", "Number 50")
    _topic("t2", "Small and big")
    topic_week_plan_repository.overwrite_assignments(
        [{"syllabus_topic_id": "t1", "week_no": 4, "position": None}]
    )

    xlsx = topic_week_import_service.build_template_xlsx("Mathematics", "Pre Year 2")
    ws = openpyxl.load_workbook(io.BytesIO(xlsx)).active
    rows = list(ws.values)

    assert list(rows[0]) == topic_week_import_service.TEMPLATE_HEADER
    body = {r[0]: r[3] for r in rows[1:]}
    assert body["t1"] == 4          # current week pehle se bhara
    assert body["t2"] in ("", None)  # bin-plan topic bhi aata hai, khali


def test_template_carries_the_title_so_rows_are_recognisable(test_db):
    _topic("t1", "Number 50")

    xlsx = topic_week_import_service.build_template_xlsx("Mathematics", "Pre Year 2")
    ws = openpyxl.load_workbook(io.BytesIO(xlsx)).active

    assert list(ws.values)[1][1] == "Number 50"


def test_template_is_scoped_to_subject_and_grade(test_db):
    _topic("t1", "Meri")
    _topic("t2", "Doosri grade", grade="Pre Year 3")

    xlsx = topic_week_import_service.build_template_xlsx("Mathematics", "Pre Year 2")
    ws = openpyxl.load_workbook(io.BytesIO(xlsx)).active

    assert [r[0] for r in list(ws.values)[1:]] == ["t1"]


def test_template_filename_slugs_the_filters(test_db):
    assert topic_week_import_service.template_filename("Mathematics", "Pre Year 2") == \
        "topic_week_plan_Mathematics_Pre_Year_2.xlsx"
    assert topic_week_import_service.template_filename(None, None) == "topic_week_plan.xlsx"


# ---- import: happy path ----


def test_import_assigns_weeks(test_db):
    _topic("t1", "Number 50")
    _topic("t2", "Small and big")

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, 3], ["t2", "Small and big", 1, 7]]), "plan.xlsx"
    )

    assert res["updated"] == 2
    assert res["errors"] == []
    assert _weeks() == {"t1": 3, "t2": 7}


def test_import_is_idempotent(test_db):
    """Dobara upload safe -- yehi SLO tagging ka waada bhi hai."""
    _topic("t1", "Number 50")
    sheet = _sheet([["t1", "Number 50", 1, 3]])

    topic_week_import_service.import_assignments(sheet, "plan.xlsx")
    res = topic_week_import_service.import_assignments(sheet, "plan.xlsx")

    assert res["updated"] == 1
    assert _weeks() == {"t1": 3}


def test_import_accepts_excel_float_style_numbers(test_db):
    """Excel 3 ko "3.0" bana deta hai -- teacher ki ghalti nahi."""
    _topic("t1", "Number 50")

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, "3.0"]]), "plan.xlsx"
    )

    assert res["updated"] == 1
    assert _weeks() == {"t1": 3}


def test_import_accepts_csv(test_db):
    _topic("t1", "Number 50")
    csv = b"syllabus_topic_id,subtopic_title,unit_no,week_no\nt1,Number 50,1,5\n"

    res = topic_week_import_service.import_assignments(csv, "plan.csv")

    assert res["updated"] == 1
    assert _weeks() == {"t1": 5}


# ---- import: clearing ----


def test_empty_week_cell_clears_the_plan(test_db):
    """Replace-set semantics: khali cell = us topic ka plan hat jaye."""
    _topic("t1", "Number 50")
    topic_week_plan_repository.overwrite_assignments(
        [{"syllabus_topic_id": "t1", "week_no": 3, "position": None}]
    )

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, ""]]), "plan.xlsx"
    )

    assert res["cleared"] == 1
    assert _weeks() == {"t1": None}


def test_week_zero_is_unassigned_not_a_clear(test_db):
    """0 ka apna matlab hai (Unassigned bucket) -- row rehti hai."""
    _topic("t1", "Number 50")

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, 0]]), "plan.xlsx"
    )

    assert res["updated"] == 1
    assert res["cleared"] == 0
    assert _weeks() == {"t1": 0}


# ---- import: errors ----


def test_unknown_topic_id_is_an_error_not_a_silent_skip(test_db):
    _topic("t1", "Number 50")

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, 3], ["ghost", "Kahin ka nahi", 1, 4]]), "plan.xlsx"
    )

    assert res["updated"] == 1
    assert len(res["errors"]) == 1
    assert "ghost" in res["errors"][0]
    assert _weeks() == {"t1": 3}


@pytest.mark.parametrize("bad", ["teen", "hafta 3", "-"])
def test_non_numeric_week_is_an_error(test_db, bad):
    _topic("t1", "Number 50")

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, bad]]), "plan.xlsx"
    )

    assert res["updated"] == 0
    assert len(res["errors"]) == 1
    assert "ginti nahi" in res["errors"][0]
    assert _weeks() == {"t1": None}


@pytest.mark.parametrize("bad", [-1, 37, 999])
def test_out_of_range_week_is_an_error(test_db, bad):
    """Hadd se bahar chup-chaap 0 likh dena teacher ko dhoka dena hoga."""
    _topic("t1", "Number 50")

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, bad]]), "plan.xlsx"
    )

    assert res["updated"] == 0
    assert "hadd se bahar" in res["errors"][0]
    assert _weeks() == {"t1": None}


def test_one_bad_row_does_not_stop_the_good_ones(test_db):
    _topic("t1", "Pehla")
    _topic("t2", "Doosra")
    _topic("t3", "Teesra")

    res = topic_week_import_service.import_assignments(
        _sheet([
            ["t1", "Pehla", 1, 2],
            ["t2", "Doosra", 1, "teen"],
            ["t3", "Teesra", 1, 4],
        ]),
        "plan.xlsx",
    )

    assert res["updated"] == 2
    assert len(res["errors"]) == 1
    assert _weeks() == {"t1": 2, "t2": None, "t3": 4}


def test_missing_required_column_is_rejected_whole(test_db):
    """Header chher diya jaye to poori file rukk jaye -- aadha plan sab se bura
    nateeja hai."""
    _topic("t1", "Number 50")

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, 3]], header=["topic", "title", "unit", "hafta"]),
        "plan.xlsx",
    )

    assert res["updated"] == 0
    assert "missing" in res["errors"][0]
    assert _weeks() == {"t1": None}


def test_empty_file_is_rejected(test_db):
    res = topic_week_import_service.import_assignments(b"", "plan.xlsx")

    assert res["updated"] == 0
    assert "khali" in res["errors"][0]


def test_unsupported_format_is_rejected(test_db):
    res = topic_week_import_service.import_assignments(b"kuch bhi", "plan.txt")

    assert res["updated"] == 0
    assert "support nahi" in res["errors"][0]


def test_corrupt_file_does_not_crash(test_db):
    res = topic_week_import_service.import_assignments(b"not really an xlsx", "plan.xlsx")

    assert res["updated"] == 0
    assert res["errors"]


def test_blank_rows_are_skipped_quietly(test_db):
    """Excel ke aakhir mein khali rows aam hain -- ye ghalti nahi."""
    _topic("t1", "Number 50")

    res = topic_week_import_service.import_assignments(
        _sheet([["t1", "Number 50", 1, 3], ["", "", "", ""]]), "plan.xlsx"
    )

    assert res["updated"] == 1
    assert res["errors"] == []


def test_round_trip_template_then_import(test_db):
    """Asal raasta: template download -> week_no bharo -> upload."""
    _topic("t1", "Number 50")
    _topic("t2", "Small and big")

    xlsx = topic_week_import_service.build_template_xlsx("Mathematics", "Pre Year 2")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        row[3].value = 6
    buf = io.BytesIO()
    wb.save(buf)

    res = topic_week_import_service.import_assignments(buf.getvalue(), "plan.xlsx")

    assert res["updated"] == 2
    assert _weeks() == {"t1": 6, "t2": 6}
